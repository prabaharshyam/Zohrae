from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

F = "Arial"
ink, muted, gold = "2A1D15", "7A6A5C", "B07A32"
f = lambda **k: Font(name=F, color=k.pop("color", ink), **k)
fill = lambda c: PatternFill("solid", fgColor=c)
HEAD, REQ, HINT, EX, INPUT = fill("2A1D15"), fill("8A5A1F"), fill("F3ECE0"), fill("ECE7E1"), fill("FFF8D6")
thin = Side(style="thin", color="D9CFC2")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")

wb = Workbook()

# ---------------- How to fill ----------------
how = wb.active
how.title = "How to fill"
how.sheet_view.showGridLines = False
how.column_dimensions["A"].width = 3
how.column_dimensions["B"].width = 28
how.column_dimensions["C"].width = 95
r = 2
how["B2"] = "Zohraé · Product sheet"
how["B2"].font = f(size=20, bold=True)
how["B3"] = "Fill in one row per product on the “Products” tab. We load everything into Shopify for you, and it appears on zohrae.in."
how["B3"].font = f(size=11, color=muted)
r = 5

def section(title):
    global r
    how.cell(r, 2, title).font = f(size=13, bold=True, color=gold)
    r += 1

def line(label, text, label_fill=None):
    global r
    a = how.cell(r, 2, label); a.font = f(bold=True, size=10); a.alignment = wrap
    if label_fill: a.fill = label_fill
    b = how.cell(r, 3, text); b.font = f(size=10); b.alignment = wrap
    how.row_dimensions[r].height = max(16, 15 * (len(text) // 95 + 1 + text.count("\n")))
    r += 1

section("1 · The colours")
line("Dark brown header", "Required column. The product can’t go live without it.", REQ)
line("Black header", "Optional. Fill in if it applies; leave empty if not.", HEAD)
line("Grey rows (3–5)", "Examples, to show the format. Leave them; we skip them.", EX)
line("Yellow cells", "Where you type. Start on row 6, one product per row.", INPUT)
line("“Status” column", "Fills itself: “✓ Ready”, or which required columns are still missing.")
for c in ("B",):
    pass
how.cell(r - 4, 2).font = f(bold=True, size=10, color="FFFFFF")
how.cell(r - 5, 2).font = f(bold=True, size=10, color="FFFFFF")
r += 1

section("2 · Filling it in")
line("Product name", "As customers should see it, e.g. “Aurelia”. Keep names unique.")
line("Category", "Pick from the list: Lamp, Vase, Planter, Art Frame, Personalised Gift, Wedding, Decor.")
line("Collections", "Optional extra shelves, separated by commas. Each product already goes into its category’s collection; add e.g. “Limited Editions” or “Gifts under ₹2,000”. The full list is on the “Collections” tab.")
line("Tagline", "One short line, under ~70 characters. Shown under the name on cards and the product page.")
line("Description", "2–4 sentences: what it is, how it feels, where it suits.")
line("Story", "Optional, in your own voice: why you made it. Shown as a quote on the product page.")
line("Price (₹)", "The starting price in rupees, numbers only, GST included (e.g. 4600). If sizes have different prices, put the smallest here.")
line("Sizes & prices", "Only if it comes in sizes. Write each size = its price, separated by semicolons:\nA3 = 1999; A2 = 3299; A1 = 4999")
line("Finishes / colours", "Only if it comes in colours or finishes, separated by semicolons. Add a price only when one costs more:\nOchre; Ivory; Sage   or   Natural oak; Matte black; Antique gold = +300")
line("Stock", "How many you have ready. Leave empty if every piece is made to order.")
line("Minimum quantity", "Only for items sold per piece (badges, magnets): the smallest order, e.g. 6. The site then shows “₹149 / piece”.")
line("Personalised?", "Yes if the customer types a name or date. Then describe it in “Personalisation details”: what’s written (name, couple’s names, date), the maximum letters, and the colours or inks available.")
line("Ready in (days)", "Making time before dispatch, e.g. 3. Shown to customers.")
line("Home page", "“Featured chapter” for the lamps that get a full-screen chapter on the home page (up to 5). “Hero” for the one signature lamp at the very top. Otherwise leave as No.")
line("Specs", "Dimensions, weight, material, finish, bulb and cable: shown in the product page’s specifications. Fill what applies.")
r += 1

section("3 · Photos")
line("One folder per product", "In Google Drive, make a folder named exactly like the “Photo folder” column (e.g. “Aurelia”) and put that product’s photos in it. Full-size originals, not WhatsApp copies.")
line("Every product", "1.jpg, 2.jpg, 3.jpg … in the order you want them shown. 1.jpg is the main photo.")
line("Lamps", "Also off.jpg and on.jpg: the same shot with the lamp switched off and on, camera on a tripod, nothing moved in between. This powers the light switch on the site. See PHOTO-GUIDE.md.")
line("Showroom (optional)", "showroom-off.jpg and showroom-on.jpg: the lamp in a dark room, landscape. Used for the big home-page chapters.")
line("Personalised pieces", "One styled photo with a real name on it, plus blank.jpg: the piece without any text, shot straight on. The site draws the customer’s name onto it.")
r += 1

section("4 · Sending it")
line("When you’re done", "Share the Drive folder (this sheet plus the photo folders) with “anyone with the link can view” and send the link. We load it into Shopify and check every product on the site with you.")
line("Prices later?", "Fine: put your best estimate now. Prices are easy to change in Shopify any time.")

# ---------------- Products ----------------
ws = wb.create_sheet("Products")
cols = [
    # (header, width, required, hint)
    ("Status", 20, False, "Fills itself"),
    ("Product name", 22, True, "e.g. Aurelia"),
    ("Category", 18, True, "Pick from the list"),
    ("Collections (extra)", 24, False, "Optional, comma-separated"),
    ("Tagline", 36, True, "One short line"),
    ("Description", 48, True, "2–4 sentences"),
    ("Story", 40, False, "Optional, your own words"),
    ("Price (₹)", 12, True, "Starting price, GST incl."),
    ("Sizes & prices", 30, False, "A3 = 1999; A2 = 3299"),
    ("Finishes / colours", 30, False, "Ochre; Ivory = +200"),
    ("Compare-at price (₹)", 14, False, "Original price, for a sale"),
    ("Stock", 9, False, "Empty = made to order"),
    ("Minimum quantity", 12, False, "Per-piece items only"),
    ("Personalised?", 13, False, "Yes / No"),
    ("Personalisation details", 40, False, "What’s written, max letters, inks"),
    ("Ready in (days)", 11, False, "Making time"),
    ("Limited edition of", 12, False, "e.g. 50, else empty"),
    ("Hours of making", 11, False, "Approx."),
    ("Dimensions", 20, False, "Ø 30 × H 38 cm"),
    ("Weight", 10, False, "1.1 kg"),
    ("Material", 22, False, ""),
    ("Finish", 22, False, "e.g. Matte, hand-sanded"),
    ("Bulb / fitting", 24, False, "Lamps only"),
    ("Cable length", 14, False, "Lamps only"),
    ("Home page", 18, False, "No / Featured chapter / Hero"),
    ("Photo folder", 18, True, "Drive folder name"),
    ("Lamp on/off photos?", 13, False, "Lamps: Yes / No"),
    ("Notes for us", 30, False, "Anything else"),
]
COL = {h: get_column_letter(i + 1) for i, (h, *_ ) in enumerate(cols)}
for i, (h, w, req, hint) in enumerate(cols, 1):
    c = ws.cell(1, i, h + (" *" if req else ""))
    c.font = f(bold=True, color="FFFFFF", size=10)
    c.fill = REQ if req else HEAD
    c.alignment = Alignment(wrap_text=True, vertical="center")
    c.border = box
    hcell = ws.cell(2, i, hint)
    hcell.font = f(italic=True, size=9, color=muted)
    hcell.fill = HINT
    hcell.alignment = wrap
    hcell.border = box
    ws.column_dimensions[get_column_letter(i)].width = w
ws.row_dimensions[1].height = 34
ws.row_dimensions[2].height = 26

examples = [
    {
        "Product name": "Aurelia", "Category": "Lamp", "Collections (extra)": "Limited Editions",
        "Tagline": "A pleated crown of gold, set on a ribbed moon.",
        "Description": "A hundred fine pleats catch the day and hold it. Switch Aurelia on and every fold turns to honey: the kind of light that makes a room slow down.",
        "Story": "I wanted the warmth of an old silk shade without the fragility.",
        "Price (₹)": 4600, "Finishes / colours": "Ochre; Ivory; Sage", "Stock": 5,
        "Personalised?": "No", "Ready in (days)": 3, "Limited edition of": 50, "Hours of making": 26,
        "Dimensions": "Ø 30 × H 38 cm", "Weight": "1.1 kg", "Material": "PLA, a plant-based bioplastic",
        "Finish": "Matte, hand-sanded", "Bulb / fitting": "E14 LED, 2700K warm white (included)",
        "Cable length": "1.8 m, braided", "Home page": "Hero", "Photo folder": "Aurelia", "Lamp on/off photos?": "Yes",
    },
    {
        "Product name": "Arches", "Category": "Art Frame",
        "Tagline": "Three terracotta arches, warm as a courtyard at four o'clock.",
        "Description": "A quiet abstract of overlapping arches in terracotta, ochre and cream.",
        "Price (₹)": 1999, "Sizes & prices": "A3 = 1999; A2 = 3299; A1 = 4999",
        "Finishes / colours": "Natural oak; Matte black; Antique gold = +300", "Personalised?": "No",
        "Ready in (days)": 5, "Material": "Giclée print on archival fine-art paper", "Home page": "No",
        "Photo folder": "Arches", "Lamp on/off photos?": "No",
    },
    {
        "Product name": "Wedding Party Badge Set", "Category": "Wedding",
        "Tagline": "The whole celebration, matched in one design.",
        "Description": "A coordinated set of role badges for the wedding party. Build it role by role and preview every badge.",
        "Price (₹)": 149, "Finishes / colours": "Ivory ribbon; Blush ribbon; Sage ribbon; Maroon ribbon",
        "Minimum quantity": 6, "Personalised?": "Yes",
        "Personalisation details": "Couple’s names (max 22 letters) + wedding date (optional). Roles: Bride, Groom, Bridesmaid, Groomsman, Cousin, Family.",
        "Ready in (days)": 10, "Home page": "No", "Photo folder": "Badge set", "Lamp on/off photos?": "No",
        "Notes for us": "Price varies with quantity: customers contact us for bulk.",
    },
]
FIRST_INPUT, LAST_INPUT = 6, 105

def status_formula(row):
    req = [("Category", "Category"), ("Tagline", "Tagline"), ("Description", "Description"), ("Price (₹)", "Price"), ("Photo folder", "Photos")]
    cond = ",".join(f'{COL[k]}{row}=""' for k, _ in req)
    missing = "&".join(f'IF({COL[k]}{row}="","{lbl} ","")' for k, lbl in req)
    b = f"{COL['Product name']}{row}"
    return f'=IF({b}="","",IF(OR({cond}),"Missing: "&{missing},"✓ Ready"))'

for n, ex in enumerate(examples):
    row = 3 + n
    for i, (h, *_ ) in enumerate(cols, 1):
        c = ws.cell(row, i, ex.get(h))
        c.font = f(italic=True, size=10, color=muted)
        c.fill = EX
        c.alignment = wrap
        c.border = box
    ws.cell(row, 1, "EXAMPLE").font = f(bold=True, size=9, color=muted)
    ws.row_dimensions[row].height = 64

for row in range(FIRST_INPUT, LAST_INPUT + 1):
    for i, (h, *_ ) in enumerate(cols, 1):
        c = ws.cell(row, i)
        c.border = box
        c.alignment = wrap
        c.font = f(size=10)
        if i > 1:
            c.fill = INPUT
    s = ws.cell(row, 1, status_formula(row))
    s.font = f(size=9, bold=True, color=gold)
    for h in ("Price (₹)", "Compare-at price (₹)"):
        ws[f"{COL[h]}{row}"].number_format = '#,##0'

ws.freeze_panes = "C3"
ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{LAST_INPUT}"

def dv(options, col, prompt):
    d = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
    d.promptTitle, d.prompt = col, prompt
    d.showInputMessage = True
    ws.add_data_validation(d)
    d.add(f"{COL[col]}{FIRST_INPUT}:{COL[col]}{LAST_INPUT}")

dv(["Lamp", "Vase", "Planter", "Art Frame", "Personalised Gift", "Wedding", "Decor"], "Category", "Pick the main category")
dv(["Yes", "No"], "Personalised?", "Does the customer type a name or date?")
dv(["Yes", "No"], "Lamp on/off photos?", "Lamps: will you send off.jpg and on.jpg?")
dv(["No", "Featured chapter", "Hero"], "Home page", "Featured chapter: up to 5 lamps. Hero: one lamp.")
for col, lo in (("Price (₹)", 1), ("Compare-at price (₹)", 1), ("Stock", 0), ("Minimum quantity", 1), ("Ready in (days)", 0), ("Limited edition of", 1), ("Hours of making", 0)):
    d = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1=str(lo), allow_blank=True)
    d.error, d.errorTitle, d.showErrorMessage = "Numbers only, please (no ₹ sign or commas).", col, True
    ws.add_data_validation(d)
    d.add(f"{COL[col]}{FIRST_INPUT}:{COL[col]}{LAST_INPUT}")

ws[f"{COL['Sizes & prices']}1"].comment = Comment("Format: Size = price; Size = price\ne.g. A3 = 1999; A2 = 3299; A1 = 4999", "Zohraé")
ws[f"{COL['Finishes / colours']}1"].comment = Comment("Separate with semicolons. Add “= +amount” only if that finish costs more.", "Zohraé")

# ---------------- Collections ----------------
cs = wb.create_sheet("Collections")
cs.sheet_view.showGridLines = False
rows = [
    ("Collection", "What goes there", "How products get in"),
    ("Lumière", "All lamps", "Automatic for Category = Lamp"),
    ("Vases & Vessels", "Vases and vessels", "Automatic for Category = Vase"),
    ("Planters", "Planters", "Automatic for Category = Planter"),
    ("Art Wall Frames", "Art frames (A3 / A2 / A1)", "Automatic for Category = Art Frame"),
    ("Objets", "Other decor", "Automatic for Category = Decor"),
    ("Personalised Gifts", "Anything with a name or date", "Automatic for Category = Personalised Gift"),
    ("Wedding Atelier", "Badges, magnets, wedding pieces", "Automatic for Category = Wedding"),
    ("Limited Editions", "Numbered pieces", "Write it in “Collections (extra)”"),
    ("Gifts under ₹2,000", "Small luxuries", "Write it in “Collections (extra)”"),
]
for i, rw in enumerate(rows, 1):
    for j, v in enumerate(rw, 1):
        c = cs.cell(i + 1, j + 1, v)
        c.border = box
        c.alignment = wrap
        c.font = f(bold=(i == 1), size=10, color="FFFFFF" if i == 1 else ink)
        if i == 1:
            c.fill = HEAD
cs.column_dimensions["A"].width = 3
cs.column_dimensions["B"].width = 24
cs.column_dimensions["C"].width = 34
cs.column_dimensions["D"].width = 40
cs.cell(len(rows) + 3, 2, "Need a new collection (e.g. “Diwali”)? Write it in “Collections (extra)” and we’ll create it.").font = f(italic=True, size=10, color=muted)

for sh in wb.worksheets:
    sh.page_setup.orientation = "landscape"
    sh.page_setup.fitToWidth = 1
    sh.page_setup.fitToHeight = 0
    sh.sheet_properties.pageSetUpPr.fitToPage = True
wb.active = 1
out = "catalogue/Zohrae-Product-Sheet.xlsx"
wb.save(out)
print(out)
